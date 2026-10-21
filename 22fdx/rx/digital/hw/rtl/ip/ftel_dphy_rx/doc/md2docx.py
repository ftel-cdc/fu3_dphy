#!/usr/bin/env python3
# ============================================================================
# ftel_dphy_rx HAS: spec.md -> DOCX in the FTEL HAS format (as HAS-CSI v1.2).
#
# Adapted from ip-dev/hw/rtl/ip/ftel_intctl/doc/md2docx.py, same strategy:
#   * The FTEL HDS template is the base: cover and logo, front matter, headers and
#     footers, section layout, theme and named styles come from it.
#   * pandoc converts the Markdown with --reference-doc=<template>, so the body
#     already uses the template's styles; the result is spliced in after the front
#     matter and replaces the template's sample sections.
#   * The conventions of HAS-CSI v1.2 (FTEL_IP_MIPI_CSI-2_CONTROLLER_HAS_v1.2.docx),
#     which is built on the same template, are applied on top:
#       - cover "<NAME> / HIGH-LEVEL ARCHITECTURE SPECIFICATION / V<x.y>", running
#         header "<DOC_ID>_V<x.y>", page number above the copyright line;
#       - figure captions below the figure in Heading 4, table captions above the
#         table in Heading 6 ("Figure 2-1. Title", "Table 3-1. Title"), centred
#         11 pt bold; Table of Tables and Table of Figures list them in TOC1 style;
#       - tables in Table Grid, fixed layout, 9 pt text, header row repeated on
#         every page, rows not split;
#       - appendices (BOS_Appendix) in the last section, Acronyms and First Review
#         at the end; revision table with a Reviewer column.
#   * A figure wider than LANDSCAPE_MIN_PX (WAVE_LANDSCAPE_MIN_PX for a timing
#     diagram, whose text is smaller) goes on a landscape page when that makes it
#     at least LANDSCAPE_GAIN larger, so its labels stay readable in print.
#   * Grounding markers of the Markdown: the warning sign becomes "(TBC)", which
#     Appendix C defines; the check mark and the globe are dropped.
#   * Contents, Table of Tables and Table of Figures are literal entries with page
#     numbers measured on a LibreOffice PDF (LibreOffice does not refresh TOC
#     fields), iterated until the page map stops moving.
# Deps: pandoc, python-docx, libreoffice, pdftotext. Figures: figures/*.png from
#   figures/render_drawio.py (block diagrams) and figures/render_wave.py (timing diagrams).
# Usage: /home/linhhk/anaconda3/bin/python3 md2docx.py   (any python3 with python-docx)
# Template: $FTEL_DPHY_RX_HAS_TEMPLATE, default
#   ~/Documents/FX1_ips/Common_Info/FTEL_CDC_IP_Hardware_Design_Specification.docx
# ============================================================================
import datetime
import io
import os
import re
import struct
import subprocess
import sys
from copy import deepcopy

try:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, Twips
    from docx.table import Table
    from docx.text.paragraph import Paragraph
except ImportError:
    sys.exit("[md2docx] ERROR: python-docx is missing; run with /home/linhhk/anaconda3/bin/python3 "
             "or install python-docx")

A_BLIP = "{http://schemas.openxmlformats.org/drawingml/2006/main}blip"
R_EMBED = "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}embed"
WP_DOCPR = "{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}docPr"
RT_IMAGE = "http://schemas.openxmlformats.org/officeDocument/2006/relationships/image"

HERE = os.path.dirname(os.path.realpath(__file__))                  # <ip>/doc
SPEC_MD = os.path.join(HERE, "spec.md")
BUILD = os.path.join(HERE, "_docx_build")                           # gitignored
TEMPLATE = os.environ.get("FTEL_DPHY_RX_HAS_TEMPLATE", os.path.expanduser(
    "~/Documents/FX1_ips/Common_Info/FTEL_CDC_IP_Hardware_Design_Specification.docx"))

DOC_ID = "FTEL_IP_MIPI_D-PHY_RX_HAS"
COVER_NAME = "MIPI D-PHY RECEIVER"
COVER_KIND = "HIGH-LEVEL ARCHITECTURE SPECIFICATION"
TITLE = "ftel_dphy_rx - MIPI D-PHY Receiver High-Level Architecture Specification"

NUM_OFFSET = 5000              # spliced list numIds stay clear of the template's
PORTRAIT_W_IN = 6.3            # text width of a portrait page is 6.47 in
PORTRAIT_H_IN = 8.5
LANDSCAPE_W_IN = 9.4           # text area of a landscape page: 9.89 x 6.67 in
LANDSCAPE_H_IN = 5.4           # leaves room for the caption
LANDSCAPE_MIN_PX = 1500        # drawio figures, text >= 18 px
WAVE_LANDSCAPE_MIN_PX = 950    # WaveDrom figures, 11 pt text: about 7 pt printed at the portrait width
LANDSCAPE_GAIN = 1.15
TEXT_W_DXA = 9304              # portrait text width used by the template's TOC tab stop
TABLE_FONT_PT = 9              # table text, as HAS-CSI
EMBED_DPI = 250                # embedded pictures: resolution at their printed size
FIGCAP, TBLCAP = "QQFIGCAPQQ", "QQTBLCAPQQ"   # plain letters: pandoc leaves them alone
TEMPLATE_LEFTOVERS = ["<IP NAME>", "FU2 SOC", "(HDS)", "[Guideline]", "[Example]", "ftel_uart",
                      "Design Review", "Quan Nguyen", "UART"]

SPEC = open(SPEC_MD, encoding="utf-8").read()
LINES = SPEC.splitlines()
FIGS, TABLES = [], []          # filled by preprocess_md(), in document order


def sh(cmd, **kw):
    return subprocess.run(cmd, check=True, **kw)


# ---------------------------------------------------------------- spec facts
def spec_version():
    m = re.search(r"\*\*DRAFT v(\d+\.\d+)\*\*", SPEC)
    if not m:
        sys.exit("[md2docx] ERROR: no '**DRAFT vX.Y**' line in spec.md")
    return m.group(1)


VERSION = spec_version()
OUT = os.path.join(HERE, f"{DOC_ID}_v{VERSION}.docx")


def split_row(line):
    body = line.strip()
    body = body[1:] if body.startswith("|") else body
    body = body[:-1] if body.endswith("|") and not body.endswith("\\|") else body
    return [c.strip().replace("\\|", "|") for c in re.split(r"(?<!\\)\|", body)]


def md_table_after(heading):
    """Rows (cells) of the first table after the line that starts with `heading`."""
    i = next((i for i, l in enumerate(LINES) if l.startswith(heading)), None)
    if i is None:
        sys.exit(f"[md2docx] ERROR: '{heading}' not found in spec.md")
    rows = []
    for line in LINES[i + 1:]:
        if line.startswith("|"):
            if not re.match(r"^\|\s*:?-", line):
                rows.append(split_row(line))
        elif rows:
            break
    return rows


WARN_RE = re.compile("\u26a0\ufe0f?")
DROP_RE = re.compile("\\s*(?:\u2705|\U0001F310)\ufe0f?")


def clean_markers(s):
    """Grounding markers: warning sign -> (TBC); check mark and globe dropped."""
    s = DROP_RE.sub("", s)
    s = WARN_RE.sub("(TBC)", s)
    return s.replace("\ufe0f", "")


def md_plain(s):
    """Markdown cell text -> plain text (front-matter and appendix tables)."""
    s = re.sub(r"(?<!!)\[([^\]]+)\]\([^)]+\)", r"\1", s)
    s = s.replace("`", "")
    s = re.sub(r"\*\*(.+?)\*\*", r"\1", s)
    return re.sub(r"\s{2,}", " ", clean_markers(s)).strip()


def revision_rows():
    """spec.md Revision and History -> template columns (Version, Date, Author, Reviewer, Description)."""
    out = []
    for ver, date, author, desc in md_table_after("## Revision and History")[1:]:
        v = re.search(r"v(\d+\.\d+)", ver).group(1)
        y, mo, d = (int(x) for x in date.split("-"))
        out.append([f"V{v}", f"{mo:02d}/{d:02d}/{y}", author, "-", md_plain(desc)])
    return out


def latest_date():
    y, mo, d = (int(x) for x in md_table_after("## Revision and History")[-1][1].split("-"))
    return datetime.date(y, mo, d)


MONTHS = ["January", "February", "March", "April", "May", "June", "July", "August", "September",
          "October", "November", "December"]


def footer_date_text():
    d = latest_date()
    return f"{MONTHS[d.month - 1]} {d.day}, {d.year}"


# ---------------------------------------------------------------- figures
def png_size(path):
    with open(path, "rb") as f:
        head = f.read(24)
    if head[:8] != b"\x89PNG\r\n\x1a\n":
        sys.exit(f"[md2docx] ERROR: {path} is not a PNG")
    return struct.unpack(">II", head[16:24])


def sentence_case(title):
    first = title.split(" ", 1)[0]
    return title[0].upper() + title[1:] if re.fullmatch(r"[a-z]+(?:-[a-z]+)*[:,]?", first) else title


def figure(num, title, stem):
    """Figure record: PNG render (figures/render_drawio.py or render_wave.py, @2x), page orientation
    and width."""
    svg, png = os.path.join(HERE, stem + ".svg"), os.path.join(HERE, stem + ".png")
    if not os.path.isfile(png):
        sys.exit(f"[md2docx] ERROR: {png} missing - run figures/render_drawio.py or render_wave.py first")
    if os.path.isfile(svg) and os.path.getmtime(png) + 60 < os.path.getmtime(svg):
        sys.exit(f"[md2docx] ERROR: {png} is older than its SVG - re-render it")
    w2, h2 = png_size(png)
    w, h = w2 / 2, h2 / 2
    aspect = w / h
    land_w = min(LANDSCAPE_W_IN, LANDSCAPE_H_IN * aspect)
    port_w = min(PORTRAIT_W_IN, PORTRAIT_H_IN * aspect)
    min_px = WAVE_LANDSCAPE_MIN_PX if "_wave_" in os.path.basename(stem) else LANDSCAPE_MIN_PX
    landscape = w >= min_px and land_w >= port_w * LANDSCAPE_GAIN
    width_in = land_w if landscape else port_w
    return {"num": num, "caption": f"Figure {num}. {sentence_case(title)}", "px": (w, h),
            "png": embed_png(png, width_in), "landscape": landscape, "width_in": width_in}


def embed_png(png, width_in):
    """Copy of the render scaled to EMBED_DPI at its printed width (the @2x renders are ~400 dpi,
    which only makes the docx heavy), in a 256-colour palette. Without Pillow the render is
    embedded as it is."""
    try:
        from PIL import Image
    except ImportError:
        return png
    os.makedirs(BUILD, exist_ok=True)
    out = os.path.join(BUILD, "fig_" + os.path.basename(png))
    im = Image.open(png)
    target = int(width_in * EMBED_DPI)
    if im.width > target:
        im = im.resize((target, round(im.height * target / im.width)), Image.LANCZOS)
    # line art: a 256-colour palette keeps the look and more than halves the size
    im.convert("RGB").quantize(colors=256, method=Image.Quantize.MEDIANCUT,
                               dither=Image.Dither.NONE).save(out, optimize=True)
    return out


# ---------------------------------------------------------------- markdown -> pandoc input
IMG_RE = re.compile(r"^!\[Figure (\d+-\d+) — (.+)\]\((figures/[^)]+)\.svg\)$")
TCAP_RE = re.compile(r"^\*\*Table ([A-Z0-9]+-\d+) — (.+)\*\*$")
HEAD_RE = re.compile(r"^(#{2,4})\s+(?:\d+(?:\.\d+)*\.?\s+)?(.*)$")
LINK_RE = re.compile(r"(?<!!)\[([^\]]+)\]\([^)]+\)")


def preprocess_md():
    """Chapters 1 to 9 and Appendices A and B (Appendix C is filled into the template's
    acronym table). Heading numbers are stripped (the template numbers them), figures point
    to their PNG with a caption paragraph, table captions become caption paragraphs, links
    become text, markers are converted, and chapter rules (---) are dropped."""
    start = next(i for i, l in enumerate(LINES) if re.match(r"##\s+1\.\s+Overview", l))
    stop = next(i for i, l in enumerate(LINES) if l.startswith("## Appendix C."))
    out = []
    for raw in LINES[start:stop]:
        line = raw.rstrip()
        if line.strip() == "---":
            continue
        m = IMG_RE.match(line.strip())
        if m:
            fig = figure(m.group(1), m.group(2), m.group(3))
            FIGS.append(fig)
            out += [f"![]({fig['png']}){{width={fig['width_in']:.2f}in}}", "", FIGCAP + fig["caption"]]
            continue
        m = TCAP_RE.match(line.strip())
        if m:
            cap = f"Table {m.group(1)}. {md_plain(m.group(2))}"
            TABLES.append(cap)
            out.append(TBLCAP + cap)
            continue
        line = clean_markers(LINK_RE.sub(r"\1", line))
        lead = re.match(r"^[ \t]*", line).group(0)          # list nesting: keep the indent
        line = lead + re.sub(r"[ \t]{2,}", " ", line[len(lead):])
        m = HEAD_RE.match(line)
        out.append(f"{m.group(1)} {m.group(2)}" if m else line)
    os.makedirs(BUILD, exist_ok=True)
    pre = os.path.join(BUILD, "_pre.md")
    open(pre, "w", encoding="utf-8").write("\n".join(out) + "\n")
    print(f"[md2docx] md     : {len(FIGS)} figures ({sum(f['landscape'] for f in FIGS)} landscape), "
          f"{len(TABLES)} table captions")
    return pre


def pandoc_body(pre_md):
    body = os.path.join(BUILD, "_body.docx")
    sh(["pandoc", pre_md, "-f", "markdown+lists_without_preceding_blankline", "-o", body, "--reference-doc", TEMPLATE,
        "--shift-heading-level-by=-1"])
    print(f"[md2docx] body   : pandoc -> {body}")
    return body


# ---------------------------------------------------------------- docx helpers
def para_style(p):
    ppr = p.find(qn("w:pPr"))
    ps = ppr.find(qn("w:pStyle")) if ppr is not None else None
    return ps.get(qn("w:val")) if ps is not None else None


def para_text(p):
    return "".join(t.text or "" for t in p.iter(qn("w:t"))).strip()


def set_text(p, text):
    """Put `text` into the first text run of paragraph element p and empty the others."""
    ts = list(p.iter(qn("w:t")))
    if not ts:
        Paragraph(p, None).add_run(text)
        return
    ts[0].text = text
    ts[0].set(qn("xml:space"), "preserve")
    for t in ts[1:]:
        t.text = ""


def strip_prefix(p, prefix):
    left = len(prefix)
    for t in p.iter(qn("w:t")):
        if left <= 0:
            break
        cut = min(left, len(t.text or ""))
        t.text = (t.text or "")[cut:]
        left -= cut


def find_para(root, text, style=None):
    for p in root.iter(qn("w:p")):
        if para_text(p) == text and (style is None or para_style(p) == style):
            return p
    return None


def merge_numbering(base, body):
    """Append body's numbering definitions to base with a collision-free id offset."""
    base_num = base.part.numbering_part.element
    body_num = body.part.numbering_part.element
    for an in body_num.findall(qn("w:abstractNum")):
        an = deepcopy(an)
        an.set(qn("w:abstractNumId"), str(int(an.get(qn("w:abstractNumId"))) + NUM_OFFSET))
        last = base_num.findall(qn("w:abstractNum"))            # schema order: abstractNum before num
        (last[-1].addnext(an) if last else base_num.insert(0, an))
    for nm in body_num.findall(qn("w:num")):
        nm = deepcopy(nm)
        nm.set(qn("w:numId"), str(int(nm.get(qn("w:numId"))) + NUM_OFFSET))
        ref = nm.find(qn("w:abstractNumId"))
        ref.set(qn("w:val"), str(int(ref.get(qn("w:val"))) + NUM_OFFSET))
        nums = base_num.findall(qn("w:num"))                     # schema order: num before
        (nums[-1].addnext(nm) if nums else base_num.append(nm))  # numIdMacAtCleanup


BOOKMARK_OFFSET = 10000


def offset_bookmarks(el):
    """pandoc numbers its heading bookmarks from 0, like the template's own: keep them unique."""
    for tag in ("w:bookmarkStart", "w:bookmarkEnd"):
        for bm in el.iter(qn(tag)):
            bm.set(qn("w:id"), str(int(bm.get(qn("w:id"))) + BOOKMARK_OFFSET))


def offset_numids(el):
    for numid in el.iter(qn("w:numId")):
        v = numid.get(qn("w:val"))
        if v and v.isdigit() and v != "0":
            numid.set(qn("w:val"), str(int(v) + NUM_OFFSET))


def merge_styles(base, body):
    """Style definitions the pandoc body uses but the template lacks."""
    have = {s.get(qn("w:styleId")) for s in base.styles.element.findall(qn("w:style"))}
    added = []
    for st in body.styles.element.findall(qn("w:style")):
        if st.get(qn("w:styleId")) not in have:
            base.styles.element.append(deepcopy(st))
            added.append(st.get(qn("w:styleId")))
    print(f"[md2docx] styles : merged {len(added)} from pandoc body")


def set_has_styles(base):
    """HAS-CSI conventions on the template styles: Heading 4 = figure caption, Heading 6 = table
    caption (centred 11 pt bold, not numbered, kept with the next paragraph); inline code in the
    body font, as HAS-CSI writes signal names."""
    for name, before, after in (("Heading 4", None, Pt(10)), ("Heading 6", Pt(12), Pt(6))):
        st = base.styles[name]
        ppr = st.element.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:numPr")) is not None:
            ppr.remove(ppr.find(qn("w:numPr")))
        pf = st.paragraph_format
        pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
        pf.keep_with_next = True
        pf.space_before, pf.space_after = before, after
        st.font.bold = True
        st.font.size = Pt(11)
    for st in base.styles.element.findall(qn("w:style")):
        if st.get(qn("w:styleId")) == "VerbatimChar":
            rpr = st.find(qn("w:rPr"))
            if rpr is not None:
                st.remove(rpr)


HEADING_NUMID = None           # the template list that numbers Heading 1-3 (set by set_heading_numbering)


def set_heading_numbering(base):
    """Heading numbers as HAS-CSI: the list is applied to each heading paragraph (numPr in the
    paragraph, see number_heading) instead of being bound to the heading styles, with the level
    format of HAS-CSI: "1." / "1.1." / "1.1.1." in Times New Roman bold 16 / 13 / 13 pt, number
    columns 0.4 / 0.55 / 0.7 in. With the template's style binding LibreOffice glues "1." to
    the title and also numbers Heading 4 / 6, which HAS-CSI uses as captions."""
    global HEADING_NUMID
    num = base.part.numbering_part.element
    HEADING_NUMID = base.styles["Heading 1"].element.find(qn("w:pPr")).find(qn("w:numPr")) \
        .find(qn("w:numId")).get(qn("w:val"))
    ref = next(n.find(qn("w:abstractNumId")).get(qn("w:val")) for n in num.findall(qn("w:num"))
               if n.get(qn("w:numId")) == HEADING_NUMID)
    an = next(a for a in num.findall(qn("w:abstractNum")) if a.get(qn("w:abstractNumId")) == ref)
    for lvl in an.findall(qn("w:lvl")):
        il = int(lvl.get(qn("w:ilvl")))
        if lvl.find(qn("w:pStyle")) is not None:
            lvl.remove(lvl.find(qn("w:pStyle")))
        if il > 2:
            continue
        width, size = {0: ("576", "32"), 1: ("792", "26"), 2: ("1008", "26")}[il]
        ppr = lvl.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            lvl.append(ppr)
        for child in list(ppr):
            ppr.remove(child)
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), width)
        ind.set(qn("w:hanging"), width)
        ppr.append(ind)
        if lvl.find(qn("w:rPr")) is not None:
            lvl.remove(lvl.find(qn("w:rPr")))
        rpr = OxmlElement("w:rPr")
        fonts = OxmlElement("w:rFonts")
        for a in ("ascii", "hAnsi", "eastAsia", "cs"):
            fonts.set(qn(f"w:{a}"), "Times New Roman")
        rpr.append(fonts)
        for tag, val in (("b", None), ("bCs", None), ("i", "0"), ("iCs", "0"), ("smallCaps", "0"),
                         ("strike", "0"), ("color", "000000"), ("sz", size), ("szCs", size),
                         ("u", "none"), ("vertAlign", "baseline")):
            el = OxmlElement(f"w:{tag}")
            if val is not None:
                el.set(qn("w:val"), val)
            rpr.append(el)
        lvl.append(rpr)
    for name in ("Heading 1", "Heading 2", "Heading 3"):
        ppr = base.styles[name].element.find(qn("w:pPr"))
        if ppr is not None and ppr.find(qn("w:numPr")) is not None:
            ppr.remove(ppr.find(qn("w:numPr")))


def number_heading(p, level):
    numpr = Paragraph(p, None)._p.get_or_add_pPr().get_or_add_numPr()
    numpr.get_or_add_ilvl().val = level
    numpr.get_or_add_numId().val = int(HEADING_NUMID)


def no_numbering(p):
    numpr = Paragraph(p, None)._p.get_or_add_pPr().get_or_add_numPr()
    numpr.get_or_add_ilvl().val = 0
    numpr.get_or_add_numId().val = 0


EMU_PER_IN = 914400


def set_picture_size(img_p, fig):
    """pandoc caps a picture at the portrait text width; set the planned size (landscape pages)."""
    w, h = fig["px"]
    cx = int(fig["width_in"] * EMU_PER_IN)
    cy = int(cx * h / w)
    for ext in img_p.iter(qn("wp:extent")):
        ext.set("cx", str(cx))
        ext.set("cy", str(cy))
    for xfrm in img_p.iter(qn("a:xfrm")):
        for ext in xfrm.findall(qn("a:ext")):
            ext.set("cx", str(cx))
            ext.set("cy", str(cy))


_docpr_id = 9000


def embed_images(el, body, base):
    """Register the pictures of a spliced element in the base package and re-target the blips."""
    global _docpr_id
    n = 0
    for blip in el.iter(A_BLIP):
        rid = blip.get(R_EMBED)
        if rid:
            new_rid, _ = base.part.get_or_add_image(io.BytesIO(body.part.rels[rid].target_part.blob))
            blip.set(R_EMBED, new_rid)
            n += 1
    for dp in el.iter(WP_DOCPR):
        _docpr_id += 1
        dp.set("id", str(_docpr_id))
    return n


def drop_orphan_images(base):
    used = {b.get(R_EMBED) for b in base.element.body.iter(A_BLIP)}
    orphans = [rid for rid, rel in list(base.part.rels.items())
               if rel.reltype == RT_IMAGE and rid not in used]
    for rid in orphans:
        base.part.drop_rel(rid)
    print(f"[md2docx] images : {len(orphans)} unused template picture(s) dropped")


# ---------------------------------------------------------------- tables
def cell_text(tc):
    return "".join(t.text or "" for t in tc.iter(qn("w:t")))


def column_widths(rows, total=TEXT_W_DXA):
    """Widths from the content: every column fits its longest word, the rest of the text
    width goes to the columns in proportion to their typical text length."""
    texts = [[cell_text(tc) for tc in tr.findall(qn("w:tc"))] for tr in rows]
    ncol = max(len(r) for r in texts)
    char, pad = 100, 230                       # dxa per character of 9 pt text, cell margins
    minw, weight = [], []
    for c in range(ncol):
        col = [r[c] for r in texts if c < len(r)]
        # a short cell (a name, a date, a value) stays on one line; a long one breaks between words
        longest = max((len(t) if len(t) <= 16 else max((len(w) for w in t.split()), default=1)
                       for t in col), default=1)
        minw.append(min(longest * char + pad, total // 3))
        lens = sorted(len(t) for t in col[1:]) or [len(col[0])]
        weight.append(max(lens[int(0.75 * (len(lens) - 1))], 0.8 * len(col[0]), 3))
    rest = total - sum(minw)
    if rest <= 0:
        return [int(m * total / sum(minw)) for m in minw]
    return [int(m + rest * w / sum(weight)) for m, w in zip(minw, weight)]


def format_table(tbl_el, doc, header=True, keep_header_align=False):
    """Table Grid, fixed layout, widths from the content, 9 pt text, header row repeated."""
    table = Table(tbl_el, doc)
    table.style = doc.styles["Table Grid"]
    table.autofit = False
    rows = tbl_el.findall(qn("w:tr"))
    widths = column_widths(rows)
    tblpr = tbl_el.find(qn("w:tblPr"))
    for old in tblpr.findall(qn("w:tblW")):
        tblpr.remove(old)
    tblw = OxmlElement("w:tblW")
    tblw.set(qn("w:w"), str(sum(widths)))
    tblw.set(qn("w:type"), "dxa")
    tblpr.insert_element_before(tblw, "w:jc", "w:tblCellSpacing", "w:tblInd", "w:tblBorders", "w:shd",
                                "w:tblLayout", "w:tblCellMar", "w:tblLook", "w:tblCaption",
                                "w:tblDescription", "w:tblPrChange")
    grid = tbl_el.find(qn("w:tblGrid"))
    for gc, w in zip(grid.findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    for i, tr in enumerate(rows):
        trpr = tr.find(qn("w:trPr"))
        if trpr is not None:
            tr.remove(trpr)
        trpr = OxmlElement("w:trPr")
        trpr.append(OxmlElement("w:cantSplit"))
        if i == 0 and header:
            trpr.append(OxmlElement("w:tblHeader"))
        tr.insert(1 if tr.find(qn("w:tblPrEx")) is not None else 0, trpr)
        for c, tc in enumerate(tr.findall(qn("w:tc"))):
            tcpr = tc.find(qn("w:tcPr"))
            if tcpr is None:
                tcpr = OxmlElement("w:tcPr")
                tc.insert(0, tcpr)
            for old in tcpr.findall(qn("w:tcW")):
                tcpr.remove(old)
            tcw = OxmlElement("w:tcW")
            tcw.set(qn("w:w"), str(widths[c] if c < len(widths) else widths[-1]))
            tcw.set(qn("w:type"), "dxa")
            tcpr.insert(0, tcw)
            for p in tc.findall(qn("w:p")):
                para = Paragraph(p, doc._body)
                para.style = None
                pf = para.paragraph_format
                pf.space_before, pf.space_after = Pt(0), Pt(0)
                if i > 0 or not header or not keep_header_align:   # template tables keep a centred header
                    pf.alignment = WD_ALIGN_PARAGRAPH.LEFT
                for run in para.runs:
                    run.font.size = Pt(TABLE_FONT_PT)
                    if i == 0 and header:
                        run.font.bold = True


def fill_simple_table(tbl_el, rows):
    """Replace the data rows of a template table (row 1 is the prototype) with `rows`."""
    trs = tbl_el.findall(qn("w:tr"))
    proto = deepcopy(trs[1])
    for tr in trs[1:]:
        tbl_el.remove(tr)
    for data in rows:
        tr = deepcopy(proto)
        tcs = tr.findall(qn("w:tc"))
        for tc, text in zip(tcs, data + [""] * (len(tcs) - len(data))):
            p = tc.find(qn("w:p"))
            for extra in tc.findall(qn("w:p"))[1:]:
                tc.remove(extra)
            set_text(p, text)
        tbl_el.append(tr)


def review_table(doc):
    """First Review table of HAS-CSI: No. | Reviewer | Comment | Response, no review yet."""
    t = doc.add_table(rows=2, cols=4)
    for c, text in enumerate(("No.", "Reviewer", "Comment", "Response")):
        t.cell(0, c).text = text
    t.cell(1, 0).text = "1"
    el = t._tbl
    el.getparent().remove(el)
    format_table(el, doc)
    widths = [700, 2100, 3252, 3252]
    for tr in el.findall(qn("w:tr")):
        for tc, w in zip(tr.findall(qn("w:tc")), widths):
            tc.find(qn("w:tcPr")).find(qn("w:tcW")).set(qn("w:w"), str(w))
    for gc, w in zip(el.find(qn("w:tblGrid")).findall(qn("w:gridCol")), widths):
        gc.set(qn("w:w"), str(w))
    return el


# ---------------------------------------------------------------- page set-up
def landscape_sections(base, sect_break, fig_paras):
    """Put each run of consecutive landscape figures (image + caption) on landscape pages:
    the paragraph before the run closes a portrait section, the last caption closes a
    landscape one. Page numbers continue; headers and footers are inherited."""
    proto = sect_break.find(qn("w:pPr")).find(qn("w:sectPr"))
    groups, cur = [], []
    for img, cap, fig in fig_paras:
        if fig["landscape"]:
            if cur and cur[-1][1].getnext() is img:
                cur.append((img, cap))
            else:
                cur = [(img, cap)]
                groups.append(cur)
        else:
            cur = []

    def close_section(p, landscape):
        sp = deepcopy(proto)
        if landscape:
            pgsz = sp.find(qn("w:pgSz"))
            w, h = pgsz.get(qn("w:w")), pgsz.get(qn("w:h"))
            pgsz.set(qn("w:w"), h)
            pgsz.set(qn("w:h"), w)
            pgsz.set(qn("w:orient"), "landscape")
        ppr = p.find(qn("w:pPr"))
        if ppr is None:
            ppr = OxmlElement("w:pPr")
            p.insert(0, ppr)
        ppr.insert_element_before(sp, "w:pPrChange")

    for g in groups:
        first_img, last_cap = g[0][0], g[-1][1]
        prev = first_img.getprevious()
        if prev is None or prev.tag != qn("w:p") or prev.find(qn("w:pPr")) is not None and \
                prev.find(qn("w:pPr")).find(qn("w:sectPr")) is not None:
            prev = OxmlElement("w:p")
            first_img.addprevious(prev)
        close_section(prev, landscape=False)
        close_section(last_cap, landscape=True)
    print(f"[md2docx] pages  : {len(groups)} landscape run(s) for "
          f"{sum(len(g) for g in groups)} figure(s)")
    return len(groups)


def set_cover(base):
    for p in base.paragraphs[:12]:
        t = p.text.strip()
        if t == "<IP NAME>":
            set_text(p._p, COVER_NAME)
        elif t.startswith("HARDWARE") and "SPECIFICATION" in t:
            set_text(p._p, COVER_KIND)
        elif re.fullmatch(r"V\d+\.\d+", t):
            set_text(p._p, f"V{VERSION}")


def set_header(base):
    """Running header of HAS-CSI: '<DOC_ID>_V<x.y>' at the left, nothing at the right."""
    n = 0
    for sec in base.sections:
        if sec.header.is_linked_to_previous:
            continue
        for p in sec.header.paragraphs:
            if "Hardware Architecture Specification" not in p.text:
                continue
            runs = p.runs
            tab = next((i for i, r in enumerate(runs) if "\t" in r.text), len(runs))
            runs[0].text = f"{DOC_ID}_V{VERSION}"
            for r in runs[1:tab] + runs[tab + 1:]:
                r.text = ""
            n += 1
    if n == 0:
        sys.exit("[md2docx] ERROR: template running header not found")


def set_footers(base):
    """Cover footer: date of the latest revision. Body footer: page number right-aligned above
    the copyright line (11 pt), as HAS-CSI."""
    txt = footer_date_text()
    date_re = re.compile(r"^[A-Z][a-z]+ \d{1,2}, \d{4}$")
    for sec in base.sections:
        for ftr in (sec.footer, sec.first_page_footer, sec.even_page_footer):
            if ftr is None or ftr.is_linked_to_previous:
                continue
            for p in ftr.paragraphs:
                if date_re.match(p.text.strip()):
                    set_text(p._p, txt)
    ftr = base.sections[1].footer._element
    paras = ftr.findall(qn("w:p"))
    copy_p = next(p for p in paras if "Copyright" in para_text(p))
    page_p = next(p for p in paras if any("PAGE" in (i.text or "") for i in p.iter(qn("w:instrText"))))
    ppr = page_p.find(qn("w:pPr"))
    if ppr is not None and ppr.find(qn("w:tabs")) is not None:
        ppr.remove(ppr.find(qn("w:tabs")))
    for r in page_p.findall(qn("w:r")):
        if r.find(qn("w:tab")) is not None and r.find(qn("w:t")) is None:
            page_p.remove(r)
    Paragraph(page_p, None).paragraph_format.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    ftr.remove(page_p)
    copy_p.addprevious(page_p)
    for r in Paragraph(copy_p, None).runs:
        r.font.size = Pt(11)
    return txt


def fill_revision_history(base):
    anchor = find_para(base.element.body, "Reversion and History")
    tbl = anchor.getnext()
    while tbl is not None and tbl.tag != qn("w:tbl"):
        tbl = tbl.getnext()
    head = tbl.findall(qn("w:tr"))[0]
    for tc in head.findall(qn("w:tc")):
        if cell_text(tc).startswith("Date"):
            set_text(tc.find(qn("w:p")), "Date")
    rows = revision_rows()
    fill_simple_table(tbl, rows)
    format_table(tbl, base, keep_header_align=True)
    print(f"[md2docx] history: {len(rows)} revision rows")


def set_properties(base):
    cp = base.core_properties
    author = md_table_after("## Revision and History")[-1][2]
    when = datetime.datetime.combine(latest_date(), datetime.time(0, 0))
    cp.title, cp.subject, cp.keywords, cp.comments = TITLE, "", "", ""
    cp.author = cp.last_modified_by = author
    cp.revision = 1
    cp.created = cp.modified = when


# ---------------------------------------------------------------- assemble
def assemble(body_docx):
    base = Document(TEMPLATE)
    body = Document(body_docx)
    bbody = base.element.body
    merge_numbering(base, body)
    merge_styles(base, body)
    set_has_styles(base)
    set_heading_numbering(base)

    # delete the template's sample sections: Heading 1 "Overview" up to the paragraph that
    # closes the body section (cover, front matter and appendix section stay)
    anchor = find_para(bbody, "Overview", "Heading1")
    breaks = [p for p in bbody.findall(qn("w:p"))
              if p.find(qn("w:pPr")) is not None and p.find(qn("w:pPr")).find(qn("w:sectPr")) is not None]
    if anchor is None or len(breaks) < 2:
        sys.exit("[md2docx] ERROR: template layout not recognised (Overview / section breaks)")
    sect_break = breaks[-1]
    kids = list(bbody)
    for el in kids[kids.index(anchor):kids.index(sect_break)]:
        bbody.remove(el)

    # splice: chapters before the section break, appendices into the appendix section
    acr_head = find_para(bbody, "Acronyms", "BOSAppendix")
    spliced, in_appx, n_img = [], False, 0
    for el in list(body.element.body):
        if el.tag == qn("w:sectPr"):
            continue
        if el.tag == qn("w:p") and para_style(el) == "Heading1" and para_text(el).startswith("Appendix "):
            in_appx = True
        el = deepcopy(el)
        offset_numids(el)
        offset_bookmarks(el)
        n_img += embed_images(el, body, base)
        (acr_head if in_appx else sect_break).addprevious(el)
        spliced.append(el)
    print(f"[md2docx] images : {n_img} figure pictures embedded")

    # pandoc paragraph styles -> Normal (HAS-CSI body text), captions, figures, appendix headings
    fig_paras, fig_i, n_tbl = [], 0, 0
    for el in spliced:
        for p in ([el] if el.tag == qn("w:p") else []):
            st, text = para_style(p), para_text(p)
            para = Paragraph(p, base._body)
            if text.startswith(FIGCAP):
                strip_prefix(p, FIGCAP)
                para.style = base.styles["Heading 4"]
                # the caption closes its figure; a heading style would chain consecutive figures
                # into one keep-together block, which then breaks between a picture and its caption
                para.paragraph_format.keep_with_next = False
                no_numbering(p)
                img = p.getprevious()
                while img is not None and img.find(".//" + qn("w:drawing")) is None:
                    img = img.getprevious()
                set_picture_size(img, FIGS[fig_i])
                fig_paras.append((img, p, FIGS[fig_i]))
                fig_i += 1
            elif text.startswith(TBLCAP):
                strip_prefix(p, TBLCAP)
                para.style = base.styles["Heading 6"]
                no_numbering(p)
            elif st == "Heading1" and text.startswith("Appendix "):
                strip_prefix(p, re.match(r"Appendix [A-Z]\.\s*", text).group(0))
                para.style = base.styles["BOS_Appendix"]
            elif st in ("Heading1", "Heading2", "Heading3"):
                number_heading(p, int(st[-1]) - 1)
            elif st in ("BodyText", "FirstParagraph", "Compact"):
                para.style = None
            if p.find(".//" + qn("w:drawing")) is not None:
                pf = para.paragraph_format
                pf.alignment = WD_ALIGN_PARAGRAPH.CENTER
                pf.keep_with_next = True
        if el.tag == qn("w:tbl"):
            format_table(el, base)
            n_tbl += 1
    if fig_i != len(FIGS):
        sys.exit(f"[md2docx] ERROR: {fig_i} figure captions spliced, {len(FIGS)} figures in the spec")
    print(f"[md2docx] tables : {n_tbl} body tables formatted")

    # appendix section: acronyms from Appendix C, review table as HAS-CSI
    acr_tbl = acr_head.getnext()
    while acr_tbl.tag != qn("w:tbl"):
        acr_tbl = acr_tbl.getnext()
    acr_rows = [[md_plain(a), md_plain(d)] for a, d in md_table_after("## Appendix C. Acronyms")[1:]]
    fill_simple_table(acr_tbl, acr_rows)
    format_table(acr_tbl, base, keep_header_align=True)
    rev_head = find_para(bbody, "Review", "BOSAppendix")
    set_text(rev_head, "First Review")
    el = rev_head.getnext()
    while el is not None and el.tag != qn("w:sectPr"):          # sample caption and table
        nxt = el.getnext()
        bbody.remove(el)
        el = nxt
    rev_tbl = review_table(base)
    rev_head.addnext(rev_tbl)
    rev_tbl.addnext(OxmlElement("w:p"))                         # a document never ends on a table
    print(f"[md2docx] appendix: {len(acr_rows)} acronyms, First Review table")

    n_land = landscape_sections(base, sect_break, fig_paras)
    set_cover(base)
    set_header(base)
    set_footers(base)
    fill_revision_history(base)
    set_properties(base)
    drop_orphan_images(base)
    base.save(OUT)
    print(f"[md2docx] saved  : {OUT}")
    return n_land


# ---------------------------------------------------------------- literal Contents / lists
FRONT_MATTER = ["Reversion and History", "Table of Tables", "Table of Figures"]


def build_pdf():
    """LibreOffice PDF of OUT: the layout oracle for the page numbers."""
    sh(["libreoffice", "--headless", "--convert-to", "pdf", "--outdir", BUILD, OUT],
       stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    return os.path.join(BUILD, os.path.splitext(os.path.basename(OUT))[0] + ".pdf")


def smart_quotes(s):
    return re.sub(r'"([^"]*)"', "\u201c\\1\u201d", s)


def _norm(s):
    s = s.replace("\u201c", '"').replace("\u201d", '"').replace("\u2019", "'")
    return re.sub(r"\s+", " ", s).strip()


def toc_targets():
    """(level, entry text, primary keys, fallback keys) of every Contents entry: the front
    matter, the chapters and sections of spec.md (numbers computed and checked against the
    Markdown), the appendices and First Review."""
    targets = [(1, t, [t], []) for t in FRONT_MATTER]
    chap = sec = 0
    appx = ord("A")
    started = False
    for line in LINES:
        m = re.match(r"^(##|###)\s+(.+?)\s*$", line)
        if not m:
            continue
        title = re.sub(r"[`*]", "", m.group(2)).strip()
        if not started:
            started = bool(re.match(r"^1\.\s+Overview$", title))
            if not started:
                continue
        ma = re.match(r"^Appendix ([A-Z])\.\s+(.*)$", title)
        if ma:
            assert ma.group(1) == chr(appx), f"appendix letter out of order: {title}"
            entry, bare = f"Appendix {chr(appx)}. {ma.group(2)}", ma.group(2)
            keys = [entry, f"Appendix {chr(appx)}.{ma.group(2)}"]
            appx += 1
            targets.append((1, smart_quotes(entry), [_norm(k)[:60] for k in keys], [_norm(bare)[:60]]))
            continue
        mn = re.match(r"^(\d+(?:\.\d+)*)\.?\s+(.*)$", title)
        if m.group(1) == "##":
            chap, sec = chap + 1, 0
            num, bare = str(chap), (mn.group(2) if mn else title)
            if mn:
                assert mn.group(1) == num, f"chapter number {mn.group(1)} != {num}"
            level = 1
        else:
            sec += 1
            num, bare = f"{chap}.{sec}", (mn.group(2) if mn else title)
            if mn:
                assert mn.group(1) == num, f"section number {mn.group(1)} != {num}"
            level = 2
        entry = f"{num}. {bare}"
        keys = [f"{num}. {bare}", f"{num}.{bare}", f"{num} {bare}"]
        targets.append((level, smart_quotes(entry), [_norm(k.split(" — ")[0])[:60] for k in keys],
                        [_norm(bare)[:60]] if len(bare) >= 8 else []))
    targets.append((1, f"Appendix {chr(appx)}. First Review",
                    [f"Appendix {chr(appx)}. First Review", f"Appendix {chr(appx)}.First Review"],
                    ["First Review"]))
    return targets


def measure_pages(pdf):
    """Page of every Contents target and every caption, searched monotonically in document
    order so that an entry of the literal lists can never shadow the real heading."""
    txt = subprocess.run(["pdftotext", pdf, "-"], check=True, capture_output=True, text=True).stdout
    pages = [_norm(p) for p in txt.split("\f")]

    def find(keys, cur):
        return next((i for i in range(cur, len(pages) + 1) if any(k in pages[i - 1] for k in keys)), None)

    toc, cur, missing = [], 1, []
    for level, entry, primary, fallback in toc_targets():
        pg = find(primary, cur) or (find(fallback, cur) if fallback else None)
        if pg is None:
            missing.append(entry)
        else:
            cur = pg
        toc.append((level, entry, pg))
    if missing:
        print(f"[md2docx] WARN: Contents entries not located in the PDF: {missing}")
        return None, None, None
    body_start = toc[len(FRONT_MATTER)][2]

    def captions(caps):
        out, c = [], body_start
        for cap in caps:
            key = _norm(cap)[:40]
            pg = find([key], c)
            if pg is None:
                print(f"[md2docx] WARN: caption not located in the PDF: {cap}")
                return None
            c = pg
            out.append((cap, pg))
        return out

    return toc, captions(TABLES), captions([f["caption"] for f in FIGS])


def _entry(body_el, style, text, page):
    p = OxmlElement("w:p")
    ppr = OxmlElement("w:pPr")
    ps = OxmlElement("w:pStyle")
    ps.set(qn("w:val"), style)
    ppr.append(ps)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:right"), "567")
    ppr.append(ind)
    p.append(ppr)
    for kind, value in (("t", text), ("tab", None), ("t", str(page))):
        r = OxmlElement("w:r")
        if kind == "tab":
            r.append(OxmlElement("w:tab"))
        else:
            t = OxmlElement("w:t")
            t.set(qn("xml:space"), "preserve")
            t.text = value
            r.append(t)
        p.append(r)
    return p


def _list_block(body_el, heading):
    """Heading paragraph of a front-matter list and the entry / field paragraphs after it."""
    paras = body_el.findall(qn("w:p"))
    idx = next((i for i, p in enumerate(paras) if para_text(p) == heading), None)
    if idx is None:
        return None, []
    entries = []
    for p in paras[idx + 1:]:
        st = (para_style(p) or "").lower()
        is_entry = st.startswith("toc") and st != "tocheading" or st.startswith("tableof") or \
            p.find(".//" + qn("w:fldChar")) is not None or p.find(".//" + qn("w:instrText")) is not None
        if is_entry:
            entries.append(p)
        else:
            break
    return paras[idx], entries


def fill_lists(toc, tables, figures):
    doc = Document(OUT)
    body_el = doc.element.body
    for heading, items, style_of in (("Contents", toc, lambda lv: f"TOC{lv}"),
                                     ("Table of Tables", tables, lambda lv: "TOC1"),
                                     ("Table of Figures", figures, lambda lv: "TOC1")):
        hd, old = _list_block(body_el, heading)
        if hd is None:
            sys.exit(f"[md2docx] ERROR: '{heading}' not found in the template front matter")
        for p in old:
            body_el.remove(p)
        anchor = hd
        for item in items:
            level, text, page = item if len(item) == 3 else (1, item[0], item[1])
            p = _entry(body_el, style_of(level), text, page)
            anchor.addnext(p)
            anchor = p
    doc.save(OUT)
    print(f"[md2docx] lists  : Contents {len(toc)}, Table of Tables {len(tables)}, "
          f"Table of Figures {len(figures)} literal entries")


# ---------------------------------------------------------------- verification
def verify_output(n_land, date_text):
    d = Document(OUT)
    body_el = d.element.body
    errors = []
    blips = body_el.findall(".//" + A_BLIP)
    dangling = [b.get(R_EMBED) for b in blips if b.get(R_EMBED) not in d.part.rels]
    if dangling or len(blips) != len(FIGS) + 1:                    # + the cover logo
        errors.append(f"{len(blips)} pictures (expected {len(FIGS)} figures + logo), dangling {dangling}")
    images = [r for r in d.part.rels.values() if r.reltype == RT_IMAGE]
    if len(images) > len(FIGS) + 1:
        errors.append(f"{len(images)} image parts for {len(FIGS) + 1} pictures")
    styles = [(para_style(p) or "") for p in body_el.iter(qn("w:p"))]
    n_toc = sum(1 for s in styles if s in ("TOC1", "TOC2"))
    n_exp = len(toc_targets()) + len(TABLES) + len(FIGS)
    if n_toc != n_exp:
        errors.append(f"{n_toc} list entries, expected {n_exp}")
    # pandoc keeps "No. 1" together with a no-break space; compare with plain spaces
    caps4 = [para_text(p).replace("\xa0", " ") for p in body_el.iter(qn("w:p")) if para_style(p) == "Heading4"]
    caps6 = [para_text(p).replace("\xa0", " ") for p in body_el.iter(qn("w:p")) if para_style(p) == "Heading6"]
    if caps4 != [f["caption"] for f in FIGS] or caps6 != TABLES:
        errors.append(f"captions: {len(caps4)} figure / {len(caps6)} table (expected {len(FIGS)} / "
                      f"{len(TABLES)})")
    land = sum(1 for sp in body_el.iter(qn("w:sectPr"))
               if sp.find(qn("w:pgSz")) is not None and sp.find(qn("w:pgSz")).get(qn("w:orient")) == "landscape")
    if land != n_land:
        errors.append(f"{land} landscape sections, expected {n_land}")
    text = "\n".join(para_text(p) for p in body_el.iter(qn("w:p")))
    hdr = [p.text for s in d.sections for h in (s.header, s.first_page_header) if not h.is_linked_to_previous
           for p in h.paragraphs]
    ftr = [p.text.strip() for s in d.sections for f in (s.footer, s.first_page_footer)
           if not f.is_linked_to_previous for p in f.paragraphs]
    left = [s for s in TEMPLATE_LEFTOVERS if s in text or any(s in h for h in hdr)]
    if left:
        errors.append(f"template text left: {left}")
    if not any(h.startswith(f"{DOC_ID}_V{VERSION}") for h in hdr):
        errors.append(f"running header is not {DOC_ID}_V{VERSION}: {hdr}")
    if date_text not in ftr:
        errors.append(f"cover footer date is not '{date_text}': {ftr}")
    markers = re.findall("[\u2705\u26a0\U0001F310\U0001F50D\ufe0f]", text)
    if markers:
        errors.append(f"{len(markers)} grounding markers left in the text")
    src = "\n".join(LINES[next(i for i, l in enumerate(LINES) if re.match(r"##\s+1\.\s+Overview", l)):
                          next(i for i, l in enumerate(LINES) if l.startswith("## Appendix C."))])
    n_warn, n_tbc = len(WARN_RE.findall(src)), text.count("(TBC)")
    if n_warn != n_tbc:
        errors.append(f"{n_warn} warning signs in the spec but {n_tbc} '(TBC)' in the docx")
    if errors:
        sys.exit("[md2docx] ERROR: " + "; ".join(errors))
    print(f"[md2docx] verify : {len(FIGS)} figures + logo, {len(TABLES)} table captions, "
          f"{n_toc} list entries, {land} landscape sections, {n_tbc} (TBC), header/footer/cover set, "
          f"no template text left")


def main():
    if not os.path.isfile(TEMPLATE):
        sys.exit(f"[md2docx] ERROR: template not found: {TEMPLATE} (set FTEL_DPHY_RX_HAS_TEMPLATE)")
    pre = preprocess_md()
    n_land = assemble(pandoc_body(pre))
    date_text = footer_date_text()
    prev = None
    for i in range(1, 6):
        cur = measure_pages(build_pdf())
        if None in cur:
            sys.exit("[md2docx] ERROR: page measurement failed (see WARN above)")
        if cur == prev:
            print(f"[md2docx] pages  : layout converged (pass {i})")
            break
        fill_lists(*cur)
        prev = cur
    else:
        sys.exit("[md2docx] ERROR: page layout did not converge in 5 passes")
    verify_output(n_land, date_text)
    older = [f for f in os.listdir(HERE) if f.startswith(DOC_ID) and f.endswith(".docx")
             and f != os.path.basename(OUT)]
    if older:
        print(f"[md2docx] NOTE: older HAS docx beside the new one: {older} (git mv / remove by hand)")


if __name__ == "__main__":
    main()
